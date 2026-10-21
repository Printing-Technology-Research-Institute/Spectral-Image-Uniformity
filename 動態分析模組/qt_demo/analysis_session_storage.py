from __future__ import annotations

import os
import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


@dataclass(frozen=True)
class SampleOutputTarget:
    """Stable file target shared by the UI and background analysis worker."""

    record_id: str
    capture_timestamp: str
    sample_dir: Path
    file_stem: str

    @property
    def uncalibrated_stem(self) -> Path:
        return self.sample_dir / f"{self.file_stem}_未校正"

    @property
    def calibrated_stem(self) -> Path:
        return self.sample_dir / f"{self.file_stem}_校正"

    @property
    def heatmap_path(self) -> Path:
        return self.sample_dir / f"{self.file_stem}_殘差熱圖.png"


class AnalysisSessionStorage:
    """Own one Start/Stop session folder and its continuously saved Excel log."""

    HEADERS = (
        "序號",
        "拍攝開始時間",
        "拍攝結束時間",
        "檔名（拍攝時間）",
        "樣本ID",
        "均勻度（%）",
        "電阻值（Ω）",
        "曝光時間（ms）",
        "Line Rate（Hz）",
        "黑校正檔",
        "白校正檔",
        "未校正HDR",
        "校正HDR",
        "殘差熱圖",
        "樣本行數",
        "分析耗時（s）",
        "狀態",
        "備註",
    )

    COLUMN_WIDTHS = (
        8,
        23,
        23,
        25,
        16,
        14,
        14,
        16,
        17,
        40,
        40,
        35,
        35,
        35,
        12,
        15,
        16,
        34,
    )

    def __init__(
        self,
        selected_root: str | Path,
        *,
        started_at: datetime,
        dark_ref_path: str | Path | None,
        white_ref_path: str | Path | None,
    ) -> None:
        self.selected_root = Path(selected_root).expanduser().resolve()
        self.selected_root.mkdir(parents=True, exist_ok=True)
        self.started_at = started_at
        self.ended_at: datetime | None = None
        self.session_timestamp = started_at.strftime("%Y%m%d_%H%M%S")
        self.session_dir = self._create_unique_directory(
            self.selected_root / f"檢測紀錄_{self.session_timestamp}"
        )
        self.workbook_path = (
            self.session_dir / f"檢測紀錄_{self.session_timestamp}.xlsx"
        )
        self.calibration_dir = self.session_dir / "校正檔"
        self.calibration_dir.mkdir(parents=True, exist_ok=True)

        self.dark_ref_original = self._resolved_text(dark_ref_path)
        self.white_ref_original = self._resolved_text(white_ref_path)
        self.dark_ref_record = self._copy_reference(dark_ref_path)
        self.white_ref_record = self._copy_reference(white_ref_path)

        self._workbook = Workbook()
        self._records = self._workbook.active
        self._records.title = "檢測紀錄"
        self._session_info = self._workbook.create_sheet("工作階段資訊")
        self._rows_by_record: dict[str, int] = {}
        self._style_workbook()
        self._write_session_info()
        self._save()

    @staticmethod
    def _resolved_text(path: str | Path | None) -> str:
        if path is None or not str(path).strip():
            return ""
        return str(Path(path).expanduser().resolve())

    @staticmethod
    def _create_unique_directory(base: Path) -> Path:
        for index in range(1000):
            candidate = base if index == 0 else base.with_name(f"{base.name}_{index:02d}")
            try:
                candidate.mkdir(parents=False, exist_ok=False)
                return candidate
            except FileExistsError:
                continue
        raise FileExistsError(f"無法建立不重複的工作階段資料夾：{base}")

    def _copy_reference(self, source: str | Path | None) -> str:
        if source is None or not str(source).strip():
            return ""
        source_path = Path(source).expanduser().resolve()
        if not source_path.is_file():
            return str(source_path)
        destination = self.calibration_dir / source_path.name
        shutil.copy2(source_path, destination)
        sidecar = source_path.with_name(source_path.name + ".json")
        if sidecar.is_file():
            shutil.copy2(sidecar, self.calibration_dir / sidecar.name)
        return str(destination.relative_to(self.session_dir))

    def _style_workbook(self) -> None:
        sheet = self._records
        sheet.append(list(self.HEADERS))
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = f"A1:R1"
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        header_border = Border(
            bottom=Side(style="medium", color="A6A6A6")
        )
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = header_border
        sheet.row_dimensions[1].height = 28
        for index, width in enumerate(self.COLUMN_WIDTHS, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width

        info = self._session_info
        info.sheet_view.showGridLines = False
        info.column_dimensions["A"].width = 25
        info.column_dimensions["B"].width = 90
        info["A1"] = "項目"
        info["B1"] = "內容"
        for cell in info[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        info.freeze_panes = "A2"

    def _write_session_info(self) -> None:
        rows = (
            ("工作階段開始時間", self.started_at),
            ("工作階段結束時間", self.ended_at),
            ("使用者選擇路徑", str(self.selected_root)),
            ("工作階段資料夾", str(self.session_dir)),
            ("原始黑校正檔", self.dark_ref_original),
            ("原始白校正檔", self.white_ref_original),
            ("工作階段黑校正檔", self.dark_ref_record),
            ("工作階段白校正檔", self.white_ref_record),
        )
        info = self._session_info
        if info.max_row > 1:
            info.delete_rows(2, info.max_row - 1)
        for label, value in rows:
            info.append([label, value])
        for row in range(2, info.max_row + 1):
            info.cell(row, 1).font = Font(bold=True, color="1F1F1F")
            info.cell(row, 1).fill = PatternFill("solid", fgColor="D9EAF7")
            info.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        info["B2"].number_format = "yyyy-mm-dd hh:mm:ss.000"
        info["B3"].number_format = "yyyy-mm-dd hh:mm:ss.000"

    def _relative(self, path: Path) -> str:
        return str(path.relative_to(self.session_dir))

    def register_sample(
        self,
        *,
        sample_id: str,
        capture_started_at: datetime,
        capture_ended_at: datetime,
        acquisition_metadata: dict[str, Any] | None,
    ) -> SampleOutputTarget:
        capture_timestamp = capture_started_at.strftime("%Y%m%d_%H%M%S_%f")
        safe_sample_id = "".join(
            ch if ch.isascii() and (ch.isalnum() or ch in "_-") else "_"
            for ch in str(sample_id).strip()
        ).strip("_-") or "sample"
        sample_dir = self._create_unique_directory(
            self.session_dir / f"{safe_sample_id}_{capture_timestamp}"
        )
        record_id = sample_dir.name
        target = SampleOutputTarget(
            record_id=record_id,
            capture_timestamp=capture_timestamp,
            sample_dir=sample_dir,
            file_stem=capture_timestamp,
        )

        metadata = acquisition_metadata if isinstance(acquisition_metadata, dict) else {}
        acquisition = metadata.get("acquisition", {})
        if not isinstance(acquisition, dict):
            acquisition = {}
        black_white = metadata.get("black_white_calibration", {})
        if not isinstance(black_white, dict):
            black_white = {}
        dark_ref = self.dark_ref_record or str(black_white.get("dark_ref_path", ""))
        white_ref = self.white_ref_record or str(black_white.get("white_ref_path", ""))

        sequence = len(self._rows_by_record) + 1
        row_values = [
            sequence,
            capture_started_at,
            capture_ended_at,
            capture_timestamp,
            sample_id,
            None,
            None,
            self._number_or_none(acquisition.get("exposure_ms")),
            self._number_or_none(acquisition.get("line_rate_hz")),
            dark_ref,
            white_ref,
            self._relative(target.uncalibrated_stem.with_suffix(".hdr")),
            self._relative(target.calibrated_stem.with_suffix(".hdr")),
            self._relative(target.heatmap_path),
            self._integer_or_none(
                metadata.get("sample_boundary", {}).get("sample_line_count")
                if isinstance(metadata.get("sample_boundary"), dict)
                else None
            ),
            None,
            "等待分析",
            "",
        ]
        self._records.append(row_values)
        row = self._records.max_row
        self._rows_by_record[record_id] = row
        self._format_record_row(row)
        self._records.auto_filter.ref = f"A1:R{row}"
        self._save()
        return target

    @staticmethod
    def _number_or_none(value: Any) -> float | None:
        if value is None or value == "":
            return None
        return float(value)

    @staticmethod
    def _integer_or_none(value: Any) -> int | None:
        if value is None or value == "":
            return None
        return int(value)

    def _format_record_row(self, row: int) -> None:
        sheet = self._records
        for column in range(1, len(self.HEADERS) + 1):
            sheet.cell(row, column).alignment = Alignment(
                vertical="center",
                wrap_text=column in (10, 11, 12, 13, 14, 18),
            )
        sheet.cell(row, 2).number_format = "yyyy-mm-dd hh:mm:ss.000"
        sheet.cell(row, 3).number_format = "yyyy-mm-dd hh:mm:ss.000"
        sheet.cell(row, 6).number_format = "0.0"
        sheet.cell(row, 7).number_format = "0.00"
        sheet.cell(row, 8).number_format = "0.000"
        sheet.cell(row, 9).number_format = "0.000"
        sheet.cell(row, 16).number_format = "0.000"
        if row % 2 == 1:
            fill = PatternFill("solid", fgColor="F3F6F9")
            for cell in sheet[row]:
                cell.fill = fill

    def complete_sample(
        self,
        record_id: str,
        *,
        uniformity: float,
        resistance_ohm: float,
        sample_lines: int,
        elapsed_seconds: float,
    ) -> None:
        row = self._rows_by_record[record_id]
        sheet = self._records
        sheet.cell(row, 6, float(uniformity))
        sheet.cell(row, 7, float(resistance_ohm))
        sheet.cell(row, 15, int(sample_lines))
        sheet.cell(row, 16, float(elapsed_seconds))
        sheet.cell(row, 17, "完成")
        sheet.cell(row, 18, "")
        self._save()

    def fail_sample(self, record_id: str, error: str) -> None:
        row = self._rows_by_record.get(record_id)
        if row is None:
            return
        detail = str(error).strip().splitlines()[-1] if str(error).strip() else "分析失敗"
        self._records.cell(row, 17, "分析失敗")
        self._records.cell(row, 18, detail[:500])
        self._save()

    def finish(self, ended_at: datetime, *, incomplete_note: str | None = None) -> None:
        self.ended_at = ended_at
        if incomplete_note:
            for record_id, row in self._rows_by_record.items():
                if self._records.cell(row, 17).value == "等待分析":
                    self._records.cell(row, 17, "未完成")
                    self._records.cell(row, 18, incomplete_note[:500])
        self._write_session_info()
        self._save()

    def _save(self) -> None:
        temporary = self.workbook_path.with_name(
            f".{self.workbook_path.stem}.tmp.xlsx"
        )
        try:
            self._workbook.save(temporary)
            os.replace(temporary, self.workbook_path)
        finally:
            temporary.unlink(missing_ok=True)

