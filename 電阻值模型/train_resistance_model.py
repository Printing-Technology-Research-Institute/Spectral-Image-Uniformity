from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
WORKSPACE_DIR = SCRIPT_DIR.parent
CAE_DIR = WORKSPACE_DIR / "CAE訓練"
if str(CAE_DIR) not in sys.path:
    sys.path.insert(0, str(CAE_DIR))

from auto_sam_mask import (  # noqa: E402
    EXCLUDED_FOR_TRAINING,
    build_sample_mask,
    choose_label_seed,
    choose_sam_mask,
    load_envi_bil,
    normalize01,
    sam_degrees,
)


@dataclass
class Sample:
    sample_id: str
    thickness_um: float
    resistance_ohm: float
    hdr_path: Path


@dataclass
class FeatureBundle:
    samples: list[Sample]
    scalar_names: list[str]
    spectral_names: list[str]
    scalar_matrix: np.ndarray
    spectral_matrix: np.ndarray
    y: np.ndarray
    rows: list[dict]


def read_resistance_excel(path: Path, data_dir: Path) -> list[Sample]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    samples: list[Sample] = []
    for row in sheet.iter_rows(values_only=True):
        for name_idx, value_idx in [(0, 1), (2, 3)]:
            if len(row) <= value_idx or row[name_idx] is None or row[value_idx] is None:
                continue
            sample_id = str(row[name_idx]).strip()
            try:
                resistance = float(row[value_idx])
            except (TypeError, ValueError):
                continue
            if sample_id.lower().endswith(".hdr"):
                sample_id = sample_id[:-4]
            thickness_text = sample_id.split("_", 1)[0].replace("u", "")
            thickness_um = float(thickness_text)
            hdr_path = data_dir / f"{sample_id}.hdr"
            samples.append(Sample(sample_id, thickness_um, resistance, hdr_path))
    if not samples:
        raise ValueError(f"No resistance rows found in {path}")
    return samples


def resolve_data_dir(user_data_dir: str | None) -> Path:
    candidates = []
    if user_data_dir:
        candidates.append(Path(user_data_dir))
    candidates.extend(
        [
            WORKSPACE_DIR / "CAE訓練" / "導電材料_RT",
            WORKSPACE_DIR / "20260528_導電材料" / "導電材料_RT",
        ]
    )
    for candidate in candidates:
        if candidate.exists():
            return candidate.resolve()
    raise FileNotFoundError("Cannot find 導電材料_RT data folder.")


def parse_sample_number(sample_id: str) -> int:
    try:
        return int(sample_id.split("_")[1])
    except Exception:
        return -1


def find_label_mask(cube: np.ndarray, args: argparse.Namespace) -> tuple[np.ndarray, dict]:
    mean_image = cube.mean(axis=2)
    sample_mask = build_sample_mask(mean_image)
    signal = normalize01(mean_image, sample_mask)
    seed, polarity, seed_info = choose_label_seed(
        signal=signal,
        sample_mask=sample_mask,
        label_polarity=args.label_polarity,
        seed_percentile=args.seed_percentile,
    )
    reference = cube[seed].mean(axis=0)
    sam = sam_degrees(cube, reference)
    label_mask, sam_threshold = choose_sam_mask(
        sam=sam,
        sample_mask=sample_mask,
        threshold_deg=args.sam_threshold_deg,
        max_label_fraction=args.max_label_fraction,
    )
    return label_mask.astype(bool), {
        "label_polarity": polarity,
        "sam_threshold_deg": float(sam_threshold),
        "seed_pixels": int(seed_info["seed_pixels"]),
        "sample_pixels": int(seed_info["sample_pixels"]),
        "label_pixels": int(label_mask.sum()),
        "label_fraction_of_sample": float(label_mask.sum() / max(sample_mask.sum(), 1)),
    }


def extract_one(sample: Sample, args: argparse.Namespace) -> tuple[dict, np.ndarray, dict]:
    cube, wavelengths, meta = load_envi_bil(sample.hdr_path, trim_bands=args.trim_bands)
    label_mask, mask_meta = find_label_mask(cube, args)
    pixels = cube[label_mask]
    if pixels.size == 0:
        raise ValueError(f"No label pixels found for {sample.sample_id}")

    mean_spectrum = pixels.mean(axis=0).astype(np.float32)
    std_spectrum = pixels.std(axis=0).astype(np.float32)
    band_axis = np.arange(mean_spectrum.shape[0], dtype=np.float32)
    slope = float(np.polyfit(band_axis, mean_spectrum, deg=1)[0])

    scalar = {
        "thickness_um": sample.thickness_um,
        "label_fraction_of_sample": mask_meta["label_fraction_of_sample"],
        "label_pixels": float(mask_meta["label_pixels"]),
        "sam_threshold_deg": mask_meta["sam_threshold_deg"],
        "mean_reflectance": float(mean_spectrum.mean()),
        "std_reflectance": float(mean_spectrum.std()),
        "roi_pixel_std_mean": float(std_spectrum.mean()),
        "spectral_range": float(np.percentile(mean_spectrum, 95) - np.percentile(mean_spectrum, 5)),
        "spectral_slope": slope,
    }
    row = {
        "sample_id": sample.sample_id,
        "sample_number": parse_sample_number(sample.sample_id),
        "thickness_um": sample.thickness_um,
        "resistance_ohm": sample.resistance_ohm,
        "hdr_path": str(sample.hdr_path),
        "hdr_lines": meta["hdr_lines"],
        "actual_lines": meta["lines"],
        "line_mismatch": bool(meta["hdr_lines"] != meta["lines"]),
        **mask_meta,
        **scalar,
    }
    return scalar, mean_spectrum, row


def build_features(samples: list[Sample], args: argparse.Namespace) -> FeatureBundle:
    scalar_rows: list[dict] = []
    spectra: list[np.ndarray] = []
    rows: list[dict] = []
    kept_samples: list[Sample] = []

    for sample in samples:
        if args.exclude_line_mismatch and sample.sample_id in EXCLUDED_FOR_TRAINING:
            print(f"[skip] {sample.sample_id}: known HDR/RAW line mismatch")
            continue
        if not sample.hdr_path.exists():
            print(f"[skip] {sample.sample_id}: missing {sample.hdr_path}")
            continue
        print(f"[feature] {sample.sample_id}")
        scalar, spectrum, row = extract_one(sample, args)
        scalar_rows.append(scalar)
        spectra.append(spectrum)
        rows.append(row)
        kept_samples.append(sample)

    if len(kept_samples) < 5:
        raise ValueError(f"Need at least 5 usable samples, got {len(kept_samples)}")

    scalar_names = list(scalar_rows[0].keys())
    scalar_matrix = np.array([[row[name] for name in scalar_names] for row in scalar_rows], dtype=np.float64)
    spectral_matrix = np.vstack(spectra).astype(np.float64)
    spectral_names = [f"spec_{i:03d}" for i in range(spectral_matrix.shape[1])]
    y = np.array([s.resistance_ohm for s in kept_samples], dtype=np.float64)
    return FeatureBundle(kept_samples, scalar_names, spectral_names, scalar_matrix, spectral_matrix, y, rows)


def safe_std(values: np.ndarray, axis=0) -> np.ndarray:
    std = values.std(axis=axis)
    return np.where(std < 1e-12, 1.0, std)


def fit_ridge(X: np.ndarray, y: np.ndarray, alpha: float) -> dict:
    x_mean = X.mean(axis=0)
    x_std = safe_std(X, axis=0)
    xs = (X - x_mean) / x_std
    y_mean = float(y.mean())
    yc = y - y_mean
    eye = np.eye(xs.shape[1])
    beta = np.linalg.solve(xs.T @ xs + alpha * eye, xs.T @ yc)
    return {"kind": "ridge", "alpha": float(alpha), "x_mean": x_mean, "x_std": x_std, "y_mean": y_mean, "beta": beta}


def predict_ridge(model: dict, X: np.ndarray) -> np.ndarray:
    xs = (X - model["x_mean"]) / model["x_std"]
    return xs @ model["beta"] + model["y_mean"]


def fit_pls(X: np.ndarray, y: np.ndarray, n_components: int) -> dict:
    x_mean = X.mean(axis=0)
    x_std = safe_std(X, axis=0)
    y_mean = float(y.mean())
    y_std = float(y.std()) if float(y.std()) > 1e-12 else 1.0
    Xh = (X - x_mean) / x_std
    yh = ((y - y_mean) / y_std).reshape(-1, 1)

    w_list, p_list, q_list = [], [], []
    max_components = min(n_components, X.shape[1], max(1, X.shape[0] - 1))
    for _ in range(max_components):
        w = (Xh.T @ yh).ravel()
        norm = np.linalg.norm(w)
        if norm < 1e-12:
            break
        w = w / norm
        t = Xh @ w
        denom = float(t.T @ t)
        if denom < 1e-12:
            break
        p = Xh.T @ t / denom
        q = float((yh.ravel() @ t) / denom)
        Xh = Xh - np.outer(t, p)
        yh = yh - t.reshape(-1, 1) * q
        w_list.append(w)
        p_list.append(p)
        q_list.append(q)

    if not w_list:
        raise ValueError("PLS failed to extract components.")
    W = np.column_stack(w_list)
    P = np.column_stack(p_list)
    Q = np.array(q_list)
    coef_scaled = W @ np.linalg.pinv(P.T @ W) @ Q
    return {
        "kind": "pls",
        "n_components": int(W.shape[1]),
        "x_mean": x_mean,
        "x_std": x_std,
        "y_mean": y_mean,
        "y_std": y_std,
        "coef_scaled": coef_scaled,
    }


def predict_pls(model: dict, X: np.ndarray) -> np.ndarray:
    xs = (X - model["x_mean"]) / model["x_std"]
    return (xs @ model["coef_scaled"]) * model["y_std"] + model["y_mean"]


def thickness_mean_loocv(samples: list[Sample], y: np.ndarray) -> np.ndarray:
    preds = np.zeros_like(y)
    thicknesses = np.array([s.thickness_um for s in samples])
    for i in range(len(y)):
        train = np.arange(len(y)) != i
        same = train & (thicknesses == thicknesses[i])
        if same.any():
            preds[i] = y[same].mean()
        else:
            preds[i] = y[train].mean()
    return preds


def loocv_predictions(X: np.ndarray, y: np.ndarray, model_kind: str, param: float | int) -> np.ndarray:
    preds = np.zeros_like(y)
    for i in range(len(y)):
        train = np.arange(len(y)) != i
        if model_kind == "ridge":
            model = fit_ridge(X[train], y[train], float(param))
            preds[i] = predict_ridge(model, X[i : i + 1])[0]
        elif model_kind == "pls":
            n_comp = min(int(param), X[train].shape[0] - 1, X[train].shape[1])
            model = fit_pls(X[train], y[train], n_comp)
            preds[i] = predict_pls(model, X[i : i + 1])[0]
        else:
            raise ValueError(model_kind)
    return preds


def metrics(y: np.ndarray, pred: np.ndarray) -> dict:
    err = pred - y
    mae = float(np.mean(np.abs(err)))
    rmse = float(np.sqrt(np.mean(err**2)))
    denom = float(np.sum((y - y.mean()) ** 2))
    r2 = float(1.0 - np.sum(err**2) / denom) if denom > 1e-12 else float("nan")
    corr = float(np.corrcoef(y, pred)[0, 1]) if len(y) > 1 and pred.std() > 1e-12 else float("nan")
    return {"mae": mae, "rmse": rmse, "r2": r2, "corr": corr}


def feature_sets(bundle: FeatureBundle) -> dict[str, tuple[np.ndarray, list[str]]]:
    return {
        "scalar": (bundle.scalar_matrix, bundle.scalar_names),
        "spectrum": (bundle.spectral_matrix, bundle.spectral_names),
        "spectrum_scalar": (
            np.hstack([bundle.spectral_matrix, bundle.scalar_matrix]),
            bundle.spectral_names + bundle.scalar_names,
        ),
    }


def evaluate_models(bundle: FeatureBundle, args: argparse.Namespace) -> tuple[list[dict], dict, np.ndarray, list[str]]:
    results: list[dict] = []
    predictions_by_key: dict[str, np.ndarray] = {}
    sets = feature_sets(bundle)

    baseline_pred = thickness_mean_loocv(bundle.samples, bundle.y)
    baseline_metrics = metrics(bundle.y, baseline_pred)
    results.append({"model": "thickness_mean_baseline", "feature_set": "thickness", "param": "", **baseline_metrics})
    predictions_by_key["thickness_mean_baseline|thickness|"] = baseline_pred

    for alpha in args.ridge_alphas:
        X, names = sets["scalar"]
        pred = loocv_predictions(X, bundle.y, "ridge", alpha)
        row = {"model": "ridge", "feature_set": "scalar", "param": alpha, **metrics(bundle.y, pred)}
        results.append(row)
        predictions_by_key[f"ridge|scalar|{alpha}"] = pred

    max_pls = min(args.max_pls_components, len(bundle.y) - 2)
    for feature_set_name in ["spectrum", "spectrum_scalar"]:
        X, names = sets[feature_set_name]
        for n_comp in range(1, max_pls + 1):
            pred = loocv_predictions(X, bundle.y, "pls", n_comp)
            row = {"model": "pls", "feature_set": feature_set_name, "param": n_comp, **metrics(bundle.y, pred)}
            results.append(row)
            predictions_by_key[f"pls|{feature_set_name}|{n_comp}"] = pred

    selectable = [row for row in results if row["model"] != "thickness_mean_baseline"]
    best = min(selectable, key=lambda r: (r["mae"], r["rmse"]))
    best_key = f"{best['model']}|{best['feature_set']}|{best['param']}"
    best_pred = predictions_by_key[best_key]
    best_X, best_names = sets[best["feature_set"]]
    return results, best, best_pred, best_names


def fit_final_model(bundle: FeatureBundle, best: dict, feature_names: list[str]) -> tuple[dict, np.ndarray]:
    X, _ = feature_sets(bundle)[best["feature_set"]]
    if best["model"] == "ridge":
        model = fit_ridge(X, bundle.y, float(best["param"]))
        pred = predict_ridge(model, X)
    elif best["model"] == "pls":
        model = fit_pls(X, bundle.y, int(best["param"]))
        pred = predict_pls(model, X)
    else:
        raise ValueError(best["model"])
    model["feature_set"] = best["feature_set"]
    model["feature_names"] = np.array(feature_names, dtype=object)
    return model, pred


def write_csv(path: Path, rows: list[dict], fieldnames: list[str] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        return
    fieldnames = fieldnames or list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def save_model(path: Path, model: dict, metadata: dict) -> None:
    arrays = {}
    for key, value in model.items():
        if isinstance(value, np.ndarray):
            arrays[key] = value
        elif isinstance(value, (float, int, str)):
            arrays[key] = np.array(value)
    np.savez(path, **arrays)
    path.with_suffix(".json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")


def plot_predictions(path: Path, y: np.ndarray, pred: np.ndarray, sample_ids: list[str], title: str) -> bool:
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return False

    path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(6.5, 5.5), dpi=160)
    ax.scatter(y, pred, s=48)
    lo = min(float(y.min()), float(pred.min())) - 0.05
    hi = max(float(y.max()), float(pred.max())) + 0.05
    ax.plot([lo, hi], [lo, hi], "k--", linewidth=1)
    for sid, x, yp in zip(sample_ids, y, pred):
        ax.annotate(sid.replace("_RT", ""), (x, yp), fontsize=7, xytext=(3, 3), textcoords="offset points")
    ax.set_xlabel("Measured resistance")
    ax.set_ylabel("Predicted resistance")
    ax.set_title(title)
    ax.grid(True, alpha=0.3)
    ax.set_xlim(lo, hi)
    ax.set_ylim(lo, hi)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    return True


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Train resistance prediction models from hyperspectral ROI features.")
    parser.add_argument("--excel", default=str(SCRIPT_DIR / "電阻值.xlsx"))
    parser.add_argument("--data-dir", default=None)
    parser.add_argument("--output-dir", default=str(SCRIPT_DIR / "resistance_model_outputs"))
    parser.add_argument("--trim-bands", type=int, default=10)
    parser.add_argument("--label-polarity", choices=["auto", "bright", "dark"], default="auto")
    parser.add_argument("--seed-percentile", type=float, default=85.0)
    parser.add_argument("--sam-threshold-deg", type=float, default=None)
    parser.add_argument("--max-label-fraction", type=float, default=0.32)
    parser.add_argument("--exclude-line-mismatch", action="store_true")
    parser.add_argument("--max-pls-components", type=int, default=8)
    parser.add_argument("--ridge-alphas", type=float, nargs="+", default=[0.01, 0.1, 1.0, 10.0, 100.0])
    return parser


def main() -> None:
    args = build_arg_parser().parse_args()
    excel_path = Path(args.excel)
    data_dir = resolve_data_dir(args.data_dir)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"[excel] {excel_path}")
    print(f"[data] {data_dir}")
    samples = read_resistance_excel(excel_path, data_dir)
    bundle = build_features(samples, args)

    feature_rows = []
    for row, spectrum in zip(bundle.rows, bundle.spectral_matrix):
        feature_rows.append({**row, **{name: float(value) for name, value in zip(bundle.spectral_names, spectrum)}})
    write_csv(output_dir / "resistance_features.csv", feature_rows)

    comparison, best, best_loocv_pred, best_feature_names = evaluate_models(bundle, args)
    comparison = sorted(comparison, key=lambda r: (r["mae"], r["rmse"]))
    write_csv(output_dir / "model_comparison.csv", comparison)

    prediction_rows = []
    for sample, measured, predicted in zip(bundle.samples, bundle.y, best_loocv_pred):
        prediction_rows.append(
            {
                "sample_id": sample.sample_id,
                "thickness_um": sample.thickness_um,
                "measured_resistance": float(measured),
                "predicted_resistance_loocv": float(predicted),
                "error": float(predicted - measured),
                "abs_error": float(abs(predicted - measured)),
            }
        )
    write_csv(output_dir / "predictions_best_loocv.csv", prediction_rows)

    final_model, full_fit_pred = fit_final_model(bundle, best, best_feature_names)
    metadata = {
        "best_model": best,
        "sample_count": len(bundle.samples),
        "feature_count": len(best_feature_names),
        "feature_names": best_feature_names,
        "excel_path": str(excel_path),
        "data_dir": str(data_dir),
        "note": "Best model selected by leave-one-out MAE among Ridge scalar and PLS spectral feature sets.",
    }
    save_model(output_dir / "resistance_model.npz", final_model, metadata)
    plot_predictions(
        output_dir / "predicted_vs_measured_loocv.png",
        bundle.y,
        best_loocv_pred,
        [s.sample_id for s in bundle.samples],
        f"LOOCV predicted vs measured ({best['model']} {best['feature_set']}, param={best['param']})",
    )

    print("[best]", best)
    print(f"[outputs] {output_dir}")


if __name__ == "__main__":
    main()
